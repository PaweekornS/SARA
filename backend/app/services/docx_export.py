"""
สร้างไฟล์ .docx ตามระเบียบสารบรรณ (FR-M5-04, 05, 06)

ฟอนต์ TH Sarabun New ขนาด 16 pt · กระดาษ A4 · ขอบตามหนังสือราชการ
ตั้งค่า w:cs (complex script) ด้วย ไม่งั้น Word จะใช้ฟอนต์อื่นเรนเดอร์ตัวอักษรไทย
"""

from __future__ import annotations

import io
import os
from datetime import date

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from app.services.thai_format import thai_date, thai_numeral

# ── Global Variables & Constants ─────────────────────────────────────────────

FONT_NAME = "TH Sarabun New"
BODY_SIZE = Pt(16)
TITLE_SIZE = Pt(20)
HEADING_SIZE = Pt(18)


# ── Helper Functions ─────────────────────────────────────────────────────────

def _new_document(template_path: str | None) -> Document:
    if template_path and os.path.exists(template_path):
        return Document(template_path)

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.0)

    style = doc.styles["Normal"]
    style.font.name = FONT_NAME
    style.font.size = BODY_SIZE
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    style.element.rPr.rFonts.set(qn("w:cs"), FONT_NAME)
    return doc


def _para(doc: Document, text: str = "", *, size=BODY_SIZE, bold=False, align=None, indent_cm=0.0):
    paragraph = doc.add_paragraph()
    if align is not None:
        paragraph.alignment = align
    if indent_cm:
        paragraph.paragraph_format.left_indent = Cm(indent_cm)
    paragraph.paragraph_format.space_after = Pt(2)

    for i, line in enumerate((text or "").split("\n")):
        run = paragraph.add_run(line)
        run.bold = bold
        run.font.size = size
        run.font.name = FONT_NAME
        run._element.rPr.rFonts.set(qn("w:cs"), FONT_NAME)
        run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
        if i < len(text.split("\n")) - 1:
            run.add_break()
    return paragraph


def _signature_block(doc: Document) -> None:
    doc.add_paragraph()
    table = doc.add_table(rows=1, cols=2)
    table.autofit = True
    left, right = table.rows[0].cells
    for cell, role in ((left, "ผู้จดรายงานการประชุม"), (right, "ผู้ตรวจรายงานการประชุม")):
        cell.text = ""
        for line in ("(ลงชื่อ) ..............................................", "", "(..............................................)", role):
            para = cell.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run(line)
            run.font.size = BODY_SIZE
            run.font.name = FONT_NAME
            run._element.rPr.rFonts.set(qn("w:cs"), FONT_NAME)


def _to_bytes(doc: Document) -> bytes:
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ── ร่างระเบียบวาระ ─────────────────────────────────────────────────────

def build_agenda_docx(
    series_name: str,
    fiscal_year: int,
    sequence_no: int,
    meeting_date: date | None,
    sections: list[tuple[int, str, list[tuple[str, str]]]] | None = None,
    items_by_section: dict[int, list[dict]] | None = None,
    template_path: str | None = None,
) -> bytes:
    doc = _new_document(template_path)

    _para(doc, "ระเบียบวาระการประชุม", size=TITLE_SIZE, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _para(doc, series_name, size=HEADING_SIZE, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _para(
        doc,
        f"ครั้งที่ {thai_numeral(sequence_no)}/{thai_numeral(fiscal_year)}",
        size=HEADING_SIZE,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    if meeting_date:
        _para(
            doc,
            f"วัน{thai_date(meeting_date)}",
            size=BODY_SIZE,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )

    _para(doc)  # บรรทัดว่าง

    section_names = {
        1: "ระเบียบวาระที่ ๑  เรื่องที่ประธานแจ้งให้ที่ประชุมทราบ",
        2: "ระเบียบวาระที่ ๒  เรื่องรับรองรายงานการประชุม",
        3: "ระเบียบวาระที่ ๓  เรื่องสืบเนื่องจากการประชุมครั้งก่อน",
        4: "ระเบียบวาระที่ ๔  เรื่องเสนอเพื่อพิจารณา",
        5: "ระเบียบวาระที่ ๕  เรื่องอื่น ๆ",
    }

    if sections is not None:
        for sec_num, sec_title, items in sections:
            header = f"ระเบียบวาระที่ {thai_numeral(sec_num)}  {sec_title}"
            _para(doc, header, bold=True)
            if not items:
                _para(doc, "- ไม่มี -", indent_cm=1.5)
                continue
            for item in items:
                if isinstance(item, tuple):
                    title, body = item
                    _para(doc, title, bold=bool(body), indent_cm=1.0)
                    if body:
                        _para(doc, body, indent_cm=1.5)
                elif isinstance(item, dict):
                    title = item.get("title", "")
                    body = item.get("body", "")
                    _para(doc, title, bold=bool(body), indent_cm=1.0)
                    if body:
                        _para(doc, body, indent_cm=1.5)
    else:
        by_sec = items_by_section or {}
        for sec_no in range(1, 6):
            _para(doc, section_names[sec_no], bold=True)
            items = by_sec.get(sec_no, [])
            if not items:
                _para(doc, "- ไม่มี -", indent_cm=1.5)
                continue
            for it in items:
                title = it.get("title", "").strip()
                body = it.get("body", "").strip()
                item_no = it.get("item_no")
                prefix = f"{sec_no}.{item_no} " if item_no else ""
                if title:
                    _para(doc, f"{prefix}{title}", bold=bool(body), indent_cm=1.0)
                if body:
                    _para(doc, body, indent_cm=1.5)

    return _to_bytes(doc)


# ── รายงานการประชุม ───────────────────────────────────────────────────

def build_minutes_docx(
    series_name: str,
    sequence_no: int,
    fiscal_year: int,
    meeting_date: date | None,
    attendees: list[str] | None = None,
    summary: str = "",
    key_points: list[str] | None = None,
    resolutions: list[dict] | None = None,
    segments: list[dict] | None = None,
    template_path: str | None = None,
) -> bytes:
    doc = _new_document(template_path)

    _para(doc, "รายงานการประชุม", size=TITLE_SIZE, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _para(doc, series_name, size=HEADING_SIZE, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _para(
        doc,
        f"ครั้งที่ {thai_numeral(sequence_no)}/{thai_numeral(fiscal_year)}",
        size=HEADING_SIZE,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    if meeting_date:
        _para(
            doc,
            f"วัน{thai_date(meeting_date)}",
            size=BODY_SIZE,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )

    _para(doc)

    _para(doc, "ผู้มาประชุม", bold=True)
    if attendees:
        for i, name in enumerate(attendees, 1):
            _para(doc, f"{thai_numeral(i)}. {name}", indent_cm=1.0)
    else:
        _para(doc, "- ตามบัญชีรายชื่อแนบท้าย -", indent_cm=1.0)

    _para(doc)
    _para(doc, "เริ่มประชุมเวลา ๐๙.๓๐ น.", indent_cm=1.0)
    _para(doc)

    if summary:
        _para(doc, "สาระสำคัญของการประชุม", bold=True)
        _para(doc, summary, indent_cm=1.0)
        _para(doc)

    if key_points:
        _para(doc, "ประเด็นที่ที่ประชุมได้หารือ", bold=True)
        for point in key_points:
            _para(doc, f"• {point}", indent_cm=1.0)
        _para(doc)

    _para(doc, "มติที่ประชุม", bold=True)
    if resolutions:
        for r in resolutions:
            ref = r.get("ref_no") or r.get("title") or ""
            text = r.get("text") or ""
            assignees = r.get("assignees") or ""
            due = r.get("due_date")
            cat = r.get("category") or ""

            head = f"{ref}: {text}" if ref and text else (ref or text)
            _para(doc, head, bold=True, indent_cm=1.0)

            details = []
            if assignees:
                details.append(f"ผู้รับผิดชอบ: {assignees}")
            if due:
                details.append(f"กำหนดเสร็จ: {thai_date(due)}")
            if cat:
                details.append(f"ประเภท: {cat}")
            if details:
                _para(doc, " · ".join(details), size=Pt(14), indent_cm=1.5)
    else:
        _para(doc, "- ไม่มีมติในการประชุมครั้งนี้ -", indent_cm=1.0)

    if segments:
        _para(doc)
        _para(doc, "บันทึกการประชุม (สรุป)", bold=True)
        for seg in segments:
            speaker = seg.get("speaker") or seg.get("speaker_name") or ""
            text = seg.get("text") or ""
            line = f"{speaker}: {text}" if speaker else text
            _para(doc, line, indent_cm=1.0)

    _para(doc)
    _para(doc, "เลิกประชุมเวลา ๑๖.๓๐ น.", indent_cm=1.0)

    _signature_block(doc)

    return _to_bytes(doc)
